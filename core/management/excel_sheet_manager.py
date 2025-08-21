from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook


def adjust_column_widths(
        worksheet: Workbook.active) -> bool:  # автонастройка ширины колонок (использовать после заполнения документа)

    for col in range(1, worksheet.max_column + 1):  # колонки "нумеруются" от 1
        max_length = 0
        column = get_column_letter(col)
        for cell in worksheet[column]:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except Exception as e:
                print(str(e))
                return False
        adjusted_width = (max_length + 2)
        worksheet.column_dimensions[column].width = adjusted_width
    return True


if __name__ == '__main__':
    wb = Workbook()
    ws = wb.active
    adjust_column_widths(ws)
